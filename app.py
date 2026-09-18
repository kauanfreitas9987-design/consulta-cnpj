import streamlit as st
import requests
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

st.set_page_config(
    page_title="Consulta de CNPJ",
    page_icon="📄",
    layout="wide"
)

st.title("📄 Consulta de CNPJ")

texto = st.text_area(
    "Cole os CNPJs (um por linha)",
    height=250
)

def gerar_excel(resultados, resumo):

    wb = Workbook()

    ws = wb.active
    ws.title = "Resultados"

    ws.append(["CNPJ", "Situação"])

    cabecalho = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    for cell in ws[1]:
        cell.fill = cabecalho
        cell.font = Font(
            color="FFFFFF",
            bold=True
        )

    for item in resultados:

        ws.append([
            item["cnpj"],
            item["situacao"]
        ])

        cor = None

        if item["situacao"] == "ATIVA":
            cor = "C6EFCE"

        elif item["situacao"] == "BAIXADA":
            cor = "FFC7CE"

        elif item["situacao"] == "INAPTA":
            cor = "FCE4D6"

        elif item["situacao"] == "SUSPENSA":
            cor = "FFF2CC"

        if cor:
            ws.cell(
                row=ws.max_row,
                column=2
            ).fill = PatternFill(
                "solid",
                fgColor=cor
            )

    resumo_ws = wb.create_sheet("Resumo")

    resumo_ws.append([
        "Situação",
        "Quantidade"
    ])

    for chave, valor in resumo.items():
        resumo_ws.append([
            chave,
            valor
        ])

    arquivo = BytesIO()

    wb.save(arquivo)

    arquivo.seek(0)

    return arquivo

if st.button("Verificar"):

    linhas = texto.strip().splitlines()

    resultados = []

    resumo = {
        "ATIVA": 0,
        "BAIXADA": 0,
        "INAPTA": 0,
        "SUSPENSA": 0,
        "ERRO": 0
    }

    barra = st.progress(0)

    for i, linha in enumerate(linhas):

        cnpj = "".join(
            filter(str.isdigit, linha)
        ).zfill(14)

        try:

            url = (
                f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}"
            )

            resposta = requests.get(
                url,
                timeout=15,
                headers={
                    "User-Agent": "Mozilla/5.0"
                }
            )

            if resposta.status_code == 200:

                dados = resposta.json()

                situacao = dados.get(
                    "descricao_situacao_cadastral",
                    "ERRO"
                ).upper()

            else:

                situacao = "ERRO"

        except:

            situacao = "ERRO"

        resultados.append({
            "cnpj": cnpj,
            "situacao": situacao
        })

        if situacao in resumo:
            resumo[situacao] += 1
        else:
            resumo["ERRO"] += 1

        barra.progress(
            (i + 1) / len(linhas)
        )

    st.success("Consulta concluída!")

    col1, col2, col3, col4, col5 = st.columns(5)

    col1.metric("ATIVA", resumo["ATIVA"])
    col2.metric("BAIXADA", resumo["BAIXADA"])
    col3.metric("INAPTA", resumo["INAPTA"])
    col4.metric("SUSPENSA", resumo["SUSPENSA"])
    col5.metric("ERRO", resumo["ERRO"])

    tabela = pd.DataFrame(resultados)

    st.dataframe(
        tabela,
        use_container_width=True
    )

    excel = gerar_excel(
        resultados,
        resumo
    )

    st.download_button(
        label="📥 Baixar Excel",
        data=excel,
        file_name="resultado_cnpjs.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
